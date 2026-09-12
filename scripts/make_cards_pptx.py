"""Сборка PPTX «Карточки в Битриксе» (Frunze Travel) — короткая версия для заказчика.

Для встречи 09.09.2026. Заказчик не технический, поэтому пять слайдов, крупные мысли,
никакой внутренней кухни: было → стало → почему → что предлагаем → итог.

Все числа — реальные замеры прода 07-09.09.2026.
Переиспользует движок стилей из make_presentation.py (как дека «Дырявое ведро»).

Запуск:  python scripts/make_cards_pptx.py
Результат: C:\\Users\\alanb\\OneDrive\\Рабочий стол\\Frunze-карточки-Битрикс-09-09.pptx
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from make_presentation import (  # noqa: E402
    Palette, SAFE_L, add_bg, add_card, add_footer, add_shape, add_text, add_title, set_cell,
)

OUT = Path(r"C:\Users\alanb\OneDrive\Рабочий стол\Frunze-карточки-Битрикс-09-09.pptx")


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# --- 1. Титул ------------------------------------------------------------------------
def slide_title(prs):
    s = _blank(prs)
    add_bg(s, dark=True)
    add_text(s, SAFE_L, 2.3, 8.0, 0.34, "FRUNZE TRAVEL · 09.09.2026", 11, True, Palette.muted_dark)
    add_text(s, SAFE_L, 2.75, 11.5, 1.0, "Карточки в Битриксе", 42, True, Palette.white)
    add_shape(s, MSO_SHAPE.RECTANGLE, SAFE_L, 4.0, 2.6, 0.07, Palette.teal, label="line")
    add_text(s, SAFE_L, 4.35, 10.5, 0.6,
             "Почему стояли на месте и что теперь работает", 17, False, Palette.muted_dark)
    add_footer(s, 1, dark=True)


# --- 2. Было / стало -----------------------------------------------------------------
def slide_result(prs):
    s = _blank(prs)
    add_bg(s)
    add_title(s, "Было и стало", kicker="Туровые карточки за месяц")

    add_card(s, SAFE_L, 2.2, 5.9, 2.5, Palette.red_soft, Palette.red)
    add_text(s, SAFE_L + 0.4, 2.45, 5.1, 0.35, "БЫЛО", 12, True, Palette.red)
    add_text(s, SAFE_L + 0.4, 2.85, 5.1, 0.9, "310", 46, True, Palette.red)
    add_text(s, SAFE_L + 0.4, 3.75, 5.1, 0.75,
             "карточек стояли в «Новом лиде» —\nэто 4 из каждых 5", 14, False, Palette.ink)

    add_card(s, 6.85, 2.2, 5.85, 2.5, Palette.green_soft, Palette.teal)
    add_text(s, 7.25, 2.45, 5.05, 0.35, "СТАЛО", 12, True, Palette.teal_dark)
    add_text(s, 7.25, 2.85, 5.05, 0.9, "0", 46, True, Palette.teal_dark)
    add_text(s, 7.25, 3.75, 5.05, 0.75,
             "карточек стоит зря — каждая\nна своей стадии", 14, False, Palette.ink)

    rows = [
        ("Стадия", "Было", "Стало"),
        ("Новый лид — никто не ответил клиенту", "310", "110"),
        ("Идёт переписка — с клиентом работают", "26", "205"),
        ("Отправили подборку туров", "31", "33"),
    ]
    table = s.shapes.add_table(len(rows), 3, Inches(SAFE_L), Inches(5.0),
                               Inches(12.1), Inches(1.55)).table
    table.columns[0].width = Inches(7.3)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(2.4)
    for r, row in enumerate(rows):
        head = r == 0
        for c, val in enumerate(row):
            fill = Palette.ink if head else Palette.white
            color = Palette.white if head else Palette.ink
            set_cell(table.cell(r, c), val, 13 if head else 14, head, color, fill,
                     PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    add_footer(s, 2)


# --- 3. Почему не работало -----------------------------------------------------------
def slide_why(prs):
    s = _blank(prs)
    add_bg(s)
    add_title(s, "Почему не работало", kicker="Одной фразой")
    add_card(s, SAFE_L, 2.2, 12.1, 1.5, Palette.grey_soft)
    add_text(s, SAFE_L + 0.5, 2.5, 11.1, 0.95,
             "Бот двигал карточку, только когда успевал выяснить у клиента всё:\n"
             "направление, даты и сколько человек едет.",
             19, True, Palette.ink)
    add_text(s, SAFE_L, 4.05, 12.1, 0.4,
             "А успевал он редко — в 8 случаях из 10 менеджер вступает в переписку раньше:",
             15, False, Palette.slate)

    cards = [
        ("85%", "диалогов забирает\nменеджер", Palette.amber),
        ("11%", "диалогов бот успевает\nвыяснить полностью", Palette.red),
        ("4 из 5", "карточек поэтому\nстояли в «Новом лиде»", Palette.red),
    ]
    x = SAFE_L
    for value, label, color in cards:
        add_card(s, x, 4.6, 3.85, 1.55)
        add_text(s, x + 0.3, 4.8, 3.25, 0.6, value, 30, True, color)
        add_text(s, x + 0.3, 5.45, 3.25, 0.6, label, 13, False, Palette.slate)
        x += 4.13
    add_text(s, SAFE_L, 6.3, 12.1, 0.3,
             "Теперь карточка едет, как только с клиентом начали работать — не дожидаясь ответов на всё.",
             14, True, Palette.teal_dark)
    add_footer(s, 3)


# --- 4. Что предлагаем ---------------------------------------------------------------
def slide_proposals(prs):
    s = _blank(prs)
    add_bg(s)
    add_title(s, "Что предлагаем дальше", kicker="От вас нужно только «да» или «нет»")
    items = [
        ("1", "Видеть, сколько денег приносит бот",
         "Продажи в CRM никто не отмечает: 17 сделок за три месяца, и все висят как новые.",
         "Бот сам спросит менеджера «клиент оплатил?» — одно нажатие в Telegram.",
         Palette.red, Palette.red_soft),
        ("2", "Считать авиабилеты отдельно",
         "Про билеты спрашивали в 94 переписках за месяц — в отчётах их не видно совсем.",
         "Будем помечать такие обращения и покажем отдельной строкой.",
         Palette.amber, Palette.amber_soft),
        ("3", "Возвращать тех, кто замолчал",
         "89 переписок без ответа, 11 из них — люди, которые уже выбирали тур.",
         "Включим мягкое напоминание на одном канале и через неделю покажем цифры.",
         Palette.teal, Palette.green_soft),
    ]
    y = 2.05
    for num, title, problem, offer, accent, soft in items:
        add_card(s, SAFE_L, y, 12.1, 1.26, soft, accent)
        add_shape(s, MSO_SHAPE.OVAL, SAFE_L + 0.32, y + 0.26, 0.44, 0.44, accent, label="num")
        add_text(s, SAFE_L + 0.32, y + 0.32, 0.44, 0.32, num, 15, True, Palette.white, PP_ALIGN.CENTER)
        add_text(s, SAFE_L + 1.0, y + 0.16, 10.8, 0.32, title, 16, True, Palette.ink)
        add_text(s, SAFE_L + 1.0, y + 0.55, 10.8, 0.3, problem, 12, False, Palette.slate)
        add_text(s, SAFE_L + 1.0, y + 0.85, 10.8, 0.3, offer, 12, True, Palette.teal_dark)
        y += 1.38
    add_text(s, SAFE_L, 6.25, 12.1, 0.3,
             "Всю работу делаем мы. Решение — за вами.", 14, True, Palette.ink)
    add_footer(s, 4)


# --- 5. Итог -------------------------------------------------------------------------
def slide_summary(prs):
    s = _blank(prs)
    add_bg(s, dark=True)
    add_text(s, SAFE_L, 1.5, 8.0, 0.3, "КОРОТКО", 11, True, Palette.muted_dark)
    add_text(s, SAFE_L, 1.9, 11.5, 0.8, "Карточки ведёт бот", 34, True, Palette.white)
    add_shape(s, MSO_SHAPE.RECTANGLE, SAFE_L, 2.85, 2.35, 0.06, Palette.teal, label="line")
    lines = [
        "Карточка уходит из «Нового лида» сама, как только с клиентом начали работать.",
        "В карточке появляется сводка: что человек хочет и когда писал в последний раз.",
        "Менеджеру не нужно ничего перетаскивать руками.",
        "Если механизм остановится — система сообщит об этом сама, а не через месяц жалоб.",
    ]
    y = 3.35
    for line in lines:
        add_shape(s, MSO_SHAPE.OVAL, SAFE_L, y + 0.12, 0.13, 0.13, Palette.teal, label="dot")
        add_text(s, SAFE_L + 0.35, y, 11.6, 0.4, line, 15.5, False, Palette.white)
        y += 0.62
    add_text(s, SAFE_L, 6.1, 12.1, 0.3,
             "Все числа — замеры боевой системы 07–09.09.2026.", 11.5, True, Palette.muted_dark)
    add_footer(s, 5, dark=True)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_title(prs)
    slide_result(prs)
    slide_why(prs)
    slide_proposals(prs)
    slide_summary(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"готово: {OUT}")


if __name__ == "__main__":
    main()
